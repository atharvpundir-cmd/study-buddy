"""Turn uploaded files into Claude content blocks.

PDFs and images are sent to Claude directly (it can read them, including diagrams and
handwriting). Word, PowerPoint and plain-text files are converted to text first.
"""

import base64
import io
import os
import zipfile

MAX_FILES = 10
MAX_TOTAL_BYTES = 20 * 1024 * 1024  # keeps the base64-encoded request under the API's 32 MB limit
MAX_TEXT_CHARS = 1_500_000  # roughly 400k tokens of extracted text across all files

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}
TEXT_TYPES = {".txt", ".md", ".csv"}
ACCEPTED = sorted([".pdf", ".docx", ".pptx", *IMAGE_TYPES, *TEXT_TYPES])

MAX_IMAGE_EDGE = 1568  # Claude's recommended longest edge; bigger images are scaled down
MAX_IMAGE_BYTES = 3_750_000  # 5 MB once base64-encoded

CONVERT_HINTS = {
    ".doc": "Please save it as .docx or PDF and try again.",
    ".ppt": "Please save it as .pptx or PDF and try again.",
    ".heic": "Please upload it as a JPG or PNG photo instead.",
    ".heif": "Please upload it as a JPG or PNG photo instead.",
    ".pages": "Please export it as PDF and try again.",
    ".key": "Please export it as PDF and try again.",
}


class FileError(ValueError):
    """A problem with an uploaded file, worded so a student can fix it."""


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


def _text_block(title: str, text: str) -> dict:
    return {"type": "document", "title": title, "source": {"type": "text", "media_type": "text/plain", "data": text}}


def _looks_like(ext: str, data: bytes) -> bool:
    """Check the file's first bytes match its extension, so a renamed file gets a clear error."""
    if ext == ".pdf":
        return data[:1024].lstrip().startswith(b"%PDF")
    if ext in (".docx", ".pptx"):
        return data[:4] == b"PK\x03\x04"
    if ext == ".png":
        return data[:8] == b"\x89PNG\r\n\x1a\n"
    if ext in (".jpg", ".jpeg"):
        return data[:3] == b"\xff\xd8\xff"
    if ext == ".gif":
        return data[:6] in (b"GIF87a", b"GIF89a")
    if ext == ".webp":
        return data[:4] == b"RIFF" and data[8:12] == b"WEBP"
    return True


def _image_block(name: str, ext: str, data: bytes) -> dict:
    media_type = IMAGE_TYPES[ext]
    try:
        from PIL import Image, ImageOps
    except ImportError:  # Pillow is optional; without it, big images are rejected instead of shrunk
        Image = None

    if Image is not None:
        try:
            img = Image.open(io.BytesIO(data))
            too_big = max(img.size) > MAX_IMAGE_EDGE or len(data) > MAX_IMAGE_BYTES
            animated = getattr(img, "is_animated", False)
            if too_big and not animated:
                img = ImageOps.exif_transpose(img)  # phone photos store rotation in EXIF
                img.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
                if img.mode not in ("RGB", "L"):
                    background = Image.new("RGB", img.size, "white")
                    background.paste(img, mask=img.convert("RGBA").split()[-1])
                    img = background
                out = io.BytesIO()
                img.save(out, format="JPEG", quality=85, optimize=True)
                data, media_type = out.getvalue(), "image/jpeg"
        except (OSError, ValueError, Image.DecompressionBombError) as exc:
            raise FileError(f"“{name}” doesn't look like a working image.") from exc

    if len(data) > MAX_IMAGE_BYTES:
        raise FileError(f"“{name}” is too big. Try a smaller photo or a screenshot.")
    return {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": _b64(data)}}


def _docx_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        lines = [f"--- Slide {number} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")
        slides.append("\n".join(lines))
    return "\n\n".join(slides)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def to_content_blocks(files):
    """Convert uploaded files into Claude content blocks.

    `files` is a list of (filename, bytes) pairs. Raises FileError with a student-friendly
    message if a file can't be used.
    """
    if len(files) > MAX_FILES:
        raise FileError(f"You can add up to {MAX_FILES} files at a time.")
    total = sum(len(data) for _, data in files)
    if total > MAX_TOTAL_BYTES:
        raise FileError(f"Your files add up to more than {MAX_TOTAL_BYTES // (1024 * 1024)} MB. Try fewer or smaller files.")

    blocks, text_chars = [], 0
    for name, data in files:
        name = os.path.basename(name or "file")
        ext = _ext(name)
        if not data:
            raise FileError(f"“{name}” is empty.")
        if ext in CONVERT_HINTS:
            raise FileError(f"StudyBuddy can't open “{name}” yet. {CONVERT_HINTS[ext]}")
        if ext not in ACCEPTED:
            raise FileError(f"StudyBuddy can't open “{name}”. Try a PDF, Word, PowerPoint, text file or photo.")
        if not _looks_like(ext, data):
            raise FileError(f"“{name}” doesn't look like a real {ext} file. Try saving it again.")

        if ext == ".pdf":
            blocks.append({"type": "document", "title": name, "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(data)}})
            continue
        if ext in IMAGE_TYPES:
            blocks.append({"type": "text", "text": f"Image: {name}"})
            blocks.append(_image_block(name, ext, data))
            continue

        try:
            if ext == ".docx":
                text = _docx_text(data)
            elif ext == ".pptx":
                text = _pptx_text(data)
            else:
                text = _decode_text(data)
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise FileError(f"StudyBuddy couldn't read “{name}”. Try saving it as PDF.") from exc
        if not text.strip():
            raise FileError(f"“{name}” doesn't have any text StudyBuddy can read. If it's scanned, try uploading it as a PDF or photo.")
        text_chars += len(text)
        if text_chars > MAX_TEXT_CHARS:
            raise FileError("Your files have too much text to read at once. Try splitting them into smaller parts.")
        blocks.append(_text_block(name, text))
    return blocks
