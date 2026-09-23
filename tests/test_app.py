import io
import json

import docx
import pytest
from PIL import Image
from pptx import Presentation

import app as app_module
from study_buddy import ai, files
from study_buddy.limits import RateLimiter


@pytest.fixture
def client():
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def ndjson(response):
    return [json.loads(line) for line in response.get_data(as_text=True).splitlines() if line.strip()]


def make_docx(text):
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Mitochondria"
    table.rows[0].cells[1].text = "Powerhouse"
    document.save(buf)
    return buf.getvalue()


def make_pptx(title, notes):
    buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = notes
    prs.save(buf)
    return buf.getvalue()


def make_png(width, height):
    buf = io.BytesIO()
    Image.new("RGBA", (width, height), (200, 100, 50, 255)).save(buf, format="PNG")
    return buf.getvalue()


# --- pages and config -------------------------------------------------------


def test_index_and_static(client):
    page = client.get("/")
    assert page.status_code == 200
    assert b"StudyBuddy" in page.data
    assert "Content-Security-Policy" in page.headers
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/../app.py").status_code == 404


def test_config(client):
    data = client.get("/api/config").get_json()
    assert data["demo"] is True
    assert [level["id"] for level in data["levels"]] == ["year6", "middle", "senior", "uni"]
    assert ".pdf" in data["accept"]


# --- streaming text modes -----------------------------------------------------


@pytest.mark.parametrize("mode", ["notes", "summary", "explain"])
def test_text_modes_stream_in_demo(client, mode):
    response = client.post("/api/generate", data={"mode": mode, "level": "year6", "text": "photosynthesis"})
    assert response.status_code == 200
    assert response.mimetype == "application/x-ndjson"
    events = ndjson(response)
    assert events[-1] == {"type": "done"}
    assert "Demo mode" in "".join(e["text"] for e in events if e["type"] == "delta")


def test_text_mode_needs_input(client):
    response = client.post("/api/generate", data={"mode": "notes", "text": "  "})
    assert response.status_code == 400
    assert "Add a file" in response.get_json()["error"]


def test_unknown_mode(client):
    assert client.post("/api/generate", data={"mode": "hack"}).status_code == 400
    assert client.post("/api/practice", data={"mode": "notes", "text": "x"}).status_code == 400


def test_text_too_long(client):
    response = client.post("/api/generate", data={"mode": "notes", "text": "a" * 50_001})
    assert response.status_code == 400


def test_generate_with_file_upload(client):
    upload = (io.BytesIO(make_docx("Cells are the building blocks of life.")), "biology.docx")
    response = client.post("/api/generate", data={"mode": "notes", "files": [upload]}, content_type="multipart/form-data")
    assert response.status_code == 200
    assert ndjson(response)[-1]["type"] == "done"


def test_bad_file_rejected_before_streaming(client):
    upload = (io.BytesIO(b"not really a pdf"), "notes.pdf")
    response = client.post("/api/generate", data={"mode": "notes", "files": [upload]}, content_type="multipart/form-data")
    assert response.status_code == 400
    assert "notes.pdf" in response.get_json()["error"]


# --- chat ---------------------------------------------------------------------


def test_chat_valid_history(client):
    history = [
        {"role": "user", "content": "What is a noun?"},
        {"role": "assistant", "content": "A naming word."},
        {"role": "user", "content": "Give me an example"},
    ]
    response = client.post("/api/generate", data={"mode": "chat", "history": json.dumps(history)})
    assert response.status_code == 200
    assert ndjson(response)[-1]["type"] == "done"


@pytest.mark.parametrize(
    "history",
    [
        "not json",
        "[]",
        json.dumps([{"role": "assistant", "content": "hi"}]),
        json.dumps([{"role": "user", "content": "hi"}, {"role": "assistant", "content": "yo"}]),
        json.dumps([{"role": "user", "content": "hi"}, {"role": "user", "content": "again"}]),
        json.dumps([{"role": "user", "content": ""}]),
        json.dumps([{"role": "user", "content": 5}]),
        json.dumps([{"role": "user", "content": "q"}, {"role": "assistant", "content": "a"}] * 20 + [{"role": "user", "content": "q"}]),
    ],
)
def test_chat_rejects_bad_history(client, history):
    response = client.post("/api/generate", data={"mode": "chat", "history": history})
    assert response.status_code == 400


# --- quizzes and flashcards -----------------------------------------------------


def test_quiz_demo(client):
    response = client.post("/api/practice", data={"mode": "quiz", "text": "plants", "count": "5"})
    assert response.status_code == 200
    quiz = response.get_json()
    assert quiz["questions"]
    for q in quiz["questions"]:
        assert 0 <= q["answer_index"] < len(q["choices"])


def test_flashcards_demo(client):
    response = client.post("/api/practice", data={"mode": "flashcards", "text": "plants"})
    assert response.status_code == 200
    assert response.get_json()["cards"][0]["front"]


def test_practice_needs_input(client):
    assert client.post("/api/practice", data={"mode": "quiz"}).status_code == 400


def test_clean_quiz_drops_broken_questions():
    data = {
        "title": " Cells ",
        "questions": [
            {"question": "Good?", "choices": ["a", "b", "c", "d"], "answer_index": 2, "explanation": "c"},
            {"question": "Bad index", "choices": ["a", "b"], "answer_index": 7, "explanation": ""},
            {"question": "", "choices": ["a", "b"], "answer_index": 0, "explanation": ""},
            {"question": "Too few", "choices": ["a"], "answer_index": 0, "explanation": ""},
        ],
    }
    cleaned = app_module._clean_quiz(data)
    assert cleaned["title"] == "Cells"
    assert [q["question"] for q in cleaned["questions"]] == ["Good?"]
    with pytest.raises(ai.AIError):
        app_module._clean_quiz({"title": "x", "questions": []})


def test_count_is_clamped(client):
    with app_module.app.test_request_context(data={"count": "999"}):
        assert app_module._count("quiz") == 25
    with app_module.app.test_request_context(data={"count": "abc"}):
        assert app_module._count("flashcards") == 12


# --- access code and rate limiting ---------------------------------------------


def test_access_code(client, monkeypatch):
    monkeypatch.setattr(app_module, "ACCESS_CODE", "owl42")
    denied = client.post("/api/practice", data={"mode": "quiz", "text": "x"})
    assert denied.status_code == 401
    assert denied.get_json()["needs_code"] is True
    wrong = client.post("/api/practice", data={"mode": "quiz", "text": "x"}, headers={"X-Access-Code": "nope"})
    assert wrong.status_code == 401
    ok = client.post("/api/practice", data={"mode": "quiz", "text": "x"}, headers={"X-Access-Code": "owl42"})
    assert ok.status_code == 200
    assert client.get("/api/config").get_json()["needsCode"] is True


def test_rate_limit(client, monkeypatch):
    monkeypatch.setattr(app_module, "limiter", RateLimiter(2, 60))
    for _ in range(2):
        assert client.post("/api/practice", data={"mode": "quiz", "text": "x"}).status_code == 200
    blocked = client.post("/api/practice", data={"mode": "quiz", "text": "x"})
    assert blocked.status_code == 429
    assert "break" in blocked.get_json()["error"]


def test_rate_limiter_spec():
    assert RateLimiter.from_spec("0").retry_after("a") == 0
    limiter = RateLimiter.from_spec("1/60")
    assert limiter.retry_after("a") == 0
    assert limiter.retry_after("a") > 0
    assert limiter.retry_after("b") == 0


def test_upload_too_large(client, monkeypatch):
    monkeypatch.setitem(app_module.app.config, "MAX_CONTENT_LENGTH", 1000)
    upload = (io.BytesIO(b"x" * 5000), "big.txt")
    response = client.post("/api/generate", data={"mode": "notes", "files": [upload]}, content_type="multipart/form-data")
    assert response.status_code == 413
    assert "too big" in response.get_json()["error"]


# --- file conversion -----------------------------------------------------------


def test_docx_to_text_block():
    [block] = files.to_content_blocks([("bio.docx", make_docx("Cells are alive."))])
    assert block["type"] == "document"
    assert block["title"] == "bio.docx"
    assert "Cells are alive." in block["source"]["data"]
    assert "Mitochondria | Powerhouse" in block["source"]["data"]


def test_pptx_includes_slides_and_notes():
    [block] = files.to_content_blocks([("deck.pptx", make_pptx("The Water Cycle", "Mention evaporation"))])
    text = block["source"]["data"]
    assert "Slide 1" in text and "The Water Cycle" in text and "Mention evaporation" in text


def test_pdf_passed_through():
    [block] = files.to_content_blocks([("a.pdf", b"%PDF-1.4 fake")])
    assert block["source"]["media_type"] == "application/pdf"


def test_text_file_encodings():
    [block] = files.to_content_blocks([("n.txt", "héllo".encode("utf-8"))])
    assert block["source"]["data"] == "héllo"


def test_small_image_kept_large_image_shrunk():
    small = files.to_content_blocks([("s.png", make_png(100, 80))])
    assert small[1]["source"]["media_type"] == "image/png"
    large = files.to_content_blocks([("l.png", make_png(4000, 3000))])
    image = large[1]
    assert image["source"]["media_type"] == "image/jpeg"
    import base64

    shrunk = Image.open(io.BytesIO(base64.b64decode(image["source"]["data"])))
    assert max(shrunk.size) <= files.MAX_IMAGE_EDGE


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("essay.doc", b"x", "save it as .docx"),
        ("photo.heic", b"x", "JPG or PNG"),
        ("virus.exe", b"MZ", "can't open"),
        ("fake.png", b"hello", "doesn't look like"),
        ("empty.txt", b"", "empty"),
        ("blank.txt", b"   ", "doesn't have any text"),
        ("broken.docx", b"PK\x03\x04garbage", "couldn't read"),
    ],
)
def test_file_errors(name, data, message):
    with pytest.raises(files.FileError, match=message):
        files.to_content_blocks([(name, data)])


def test_too_many_files():
    with pytest.raises(files.FileError, match="up to"):
        files.to_content_blocks([("a.txt", b"a")] * (files.MAX_FILES + 1))


# --- request sent to Claude --------------------------------------------------------


class FakeStream:
    def __init__(self, texts, final):
        self._texts, self._final = texts, final

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    @property
    def text_stream(self):
        return iter(self._texts)

    def get_final_message(self):
        return self._final


class FakeBlock:
    def __init__(self, text):
        self.type, self.text = "text", text


class FakeMessage:
    def __init__(self, stop_reason, content=()):
        self.stop_reason, self.content = stop_reason, list(content)


class FakeClient:
    def __init__(self, stream):
        self.calls = []
        outer = self

        class Messages:
            def stream(self, **params):
                outer.calls.append(params)
                return stream

        class Beta:
            messages = Messages()

        self.beta = Beta()


@pytest.fixture
def live(monkeypatch):
    monkeypatch.setenv("STUDY_BUDDY_DEMO", "0")

    def install(stream):
        fake = FakeClient(stream)
        monkeypatch.setattr(ai, "_get_client", lambda: fake)
        return fake

    return install


def test_stream_text_request_shape(live):
    fake = live(FakeStream(["Hel", "lo"], FakeMessage("end_turn")))
    out = "".join(ai.stream_text("sys", [{"role": "user", "content": "hi"}], "medium"))
    assert out == "Hello"
    params = fake.calls[0]
    assert params["model"] == ai.MODEL
    assert params["thinking"] == {"type": "adaptive"}
    assert params["output_config"] == {"effort": "medium"}
    assert params["fallbacks"] == "default"
    assert params["betas"] == [ai.FALLBACK_BETA]
    assert params["cache_control"] == {"type": "ephemeral"}


def test_stream_text_refusal_and_truncation(live):
    live(FakeStream(["partial"], FakeMessage("refusal")))
    with pytest.raises(ai.AIError, match="can't help"):
        list(ai.stream_text("sys", [], "high"))
    live(FakeStream(["long"], FakeMessage("max_tokens")))
    assert "carry on" in "".join(ai.stream_text("sys", [], "high"))


def test_generate_json_uses_schema(live):
    payload = {"title": "T", "cards": [{"front": "a", "back": "b"}]}
    fake = live(FakeStream([], FakeMessage("end_turn", [FakeBlock(json.dumps(payload))])))
    schema = {"type": "object", "properties": {"cards": {}}}
    assert ai.generate_json("sys", [], "high", schema) == payload
    assert fake.calls[0]["output_config"]["format"] == {"type": "json_schema", "schema": schema}


def test_generate_json_errors(live):
    live(FakeStream([], FakeMessage("end_turn", [FakeBlock("{not json")])))
    with pytest.raises(ai.AIError):
        ai.generate_json("sys", [], "high", {"properties": {}})
    live(FakeStream([], FakeMessage("max_tokens")))
    with pytest.raises(ai.AIError, match="fewer"):
        ai.generate_json("sys", [], "high", {"properties": {}})


def test_api_errors_become_friendly_messages(live, client):
    import anthropic
    import httpx2

    class Exploding(FakeStream):
        def __enter__(self):
            request = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
            raise anthropic.RateLimitError("slow down", response=httpx2.Response(429, request=request), body=None)

    live(Exploding([], None))
    events = ndjson(client.post("/api/generate", data={"mode": "notes", "text": "x"}))
    assert events[-1]["type"] == "error"
    assert "Lots of people" in events[-1]["message"]
    response = client.post("/api/practice", data={"mode": "quiz", "text": "x"})
    assert response.status_code == 502
    assert "Lots of people" in response.get_json()["error"]


def test_files_reach_claude_before_the_instruction(live, client):
    fake = live(FakeStream(["ok"], FakeMessage("end_turn")))
    upload = (io.BytesIO(b"%PDF-1.4 fake"), "chapter.pdf")
    client.post("/api/generate", data={"mode": "notes", "level": "uni", "text": "focus on ch 2", "files": [upload]}, content_type="multipart/form-data").get_data()
    params = fake.calls[0]
    content = params["messages"][0]["content"]
    assert content[0]["type"] == "document"
    assert content[-1]["type"] == "text" and "focus on ch 2" in content[-1]["text"]
    assert "University" in params["system"]
